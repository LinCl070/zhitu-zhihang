"""Fill the official 作品设计方案 template for the governance-agent submission.

All simulation numbers are read from docs/deliverables/simulation-sensitivity.json,
which backend/scripts/simulation_sensitivity.py generates; proposal_content.py only
formats them. The same blocks also render the markdown copy. Run from repo root:

    python tools/build_proposal_docx.py
"""

import copy
import json
import re
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402
from matplotlib.patches import FancyBboxPatch  # noqa: E402

TOOLS = Path(__file__).resolve().parent
ROOT = TOOLS.parent
sys.path.insert(0, str(TOOLS))
from proposal_content import (  # noqa: E402
    COVER,
    LIMITS,
    build_blocks,
    build_registration,
    char_count,
)

TEMPLATE = ROOT / "中国仿真创新应用大赛作品方案模板.docx"
OUTPUT = ROOT / "中国仿真创新应用大赛作品方案-职途智航.docx"
REG_TEMPLATE = ROOT / "中国仿真创新应用大赛报名表.docx"
REG_OUTPUT = ROOT / "中国仿真创新应用大赛报名表-职途智航.docx"
REG_MARKDOWN = ROOT / "docs" / "deliverables" / "05-报名表文案.md"
MARKDOWN = ROOT / "docs" / "deliverables" / "04-仿真作品方案书.md"
DATA = ROOT / "docs" / "deliverables" / "simulation-sensitivity.json"
FIG_DIR = ROOT / "docs" / "deliverables" / "figures"

LABELS = {
    "baseline": "现状基线",
    "skill_training": "技能补强培训",
    "precision_coaching": "精准帮扶",
    "job_push": "岗位精准推送",
}
# Marker + line style per scenario so curves stay distinguishable in grayscale print.
STYLES = {
    "baseline": ("o", "-", "#4C4C4C"),
    "skill_training": ("s", "--", "#1F77B4"),
    "precision_coaching": ("^", "-.", "#D62728"),
    "job_push": ("D", ":", "#2CA02C"),
}
plt.rcParams["font.sans-serif"] = ["SimHei", "Microsoft YaHei"]
plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["savefig.dpi"] = 200


def _box(ax, x, y, w, h, title, sub="", fill="#EAF1F8"):
    ax.add_patch(
        FancyBboxPatch(
            (x - w / 2, y - h / 2), w, h, boxstyle="round,pad=0.02,rounding_size=0.06",
            linewidth=1.2, edgecolor="#1F4E79", facecolor=fill,
        )
    )
    ax.text(x, y + (0.15 if sub else 0), title, ha="center", va="center", fontsize=11,
            fontweight="bold", color="#17324D")
    if sub:
        ax.text(x, y - 0.17, sub, ha="center", va="center", fontsize=8, color="#333333",
                linespacing=1.3)


def _arrow(ax, start, end, text="", rad=0.0):
    ax.annotate("", xy=end, xytext=start, arrowprops=dict(
        arrowstyle="-|>", color="#1F4E79", lw=1.3, connectionstyle=f"arc3,rad={rad}"))
    if text:
        mx, my = (start[0] + end[0]) / 2, (start[1] + end[1]) / 2
        ax.text(mx, my + 0.1, text, ha="center", fontsize=8, color="#5B6573")


def fig_architecture(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(8.2, 4.4))
    ax.set_xlim(0, 9)
    ax.set_ylim(0, 4.6)
    ax.axis("off")
    top = [(1.5, "感知", "群体画像聚合\nk=5 小样本抑制"), (4.5, "预警", "六因子风险模型\n教师干预清单"),
           (7.5, "仿真推演", "ABM 多策略实验\n共同随机数 + CI")]
    bottom = [(7.5, "决策建议", "确定性策略报告\nAI 解读需人工审核"),
              (4.5, "人工确认", "管理员采纳\n写入审计，不自动执行"),
              (1.5, "反馈", "教师指导意见\n回流风险看板")]
    for x, t, s in top:
        _box(ax, x, 3.6, 2.3, 1.0, t, s)
    for x, t, s in bottom:
        _box(ax, x, 1.9, 2.3, 1.0, t, s)
    _arrow(ax, (2.65, 3.6), (3.35, 3.6))
    _arrow(ax, (5.65, 3.6), (6.35, 3.6))
    _arrow(ax, (7.5, 3.1), (7.5, 2.4))
    _arrow(ax, (6.35, 1.9), (5.65, 1.9))
    _arrow(ax, (3.35, 1.9), (2.65, 1.9))
    _arrow(ax, (1.5, 2.4), (1.5, 3.1))
    _box(ax, 4.5, 0.55, 8.4, 0.75, "数据与模型底座",
         "脱敏画像 · 行动计划 · 已发布岗位 · 合成群体 · 岗位匹配与风险模型（业务与仿真共用）",
         fill="#F3F5F7")
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
def fig_weekly_flow(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(8.6, 3.0))
    ax.set_xlim(0, 10.4)
    ax.set_ylim(0, 3.2)
    ax.axis("off")
    steps = [
        (1.25, "① 教师服务", "会谈数=教师×每周会谈\n基线随机/精准按风险"),
        (3.9, "② 学生行动", "按主动性概率行动\n探索→准备→求职"),
        (6.55, "③ 投递匹配", "可见岗位前 3 个\n≥60 分且有容量"),
        (9.2, "④ 时间推进", "距毕业减 0.25 月\n刷新匹配与风险"),
    ]
    for x, t, s in steps:
        _box(ax, x, 2.2, 2.25, 1.2, t, s)
    for i in range(3):
        _arrow(ax, (steps[i][0] + 1.15, 2.2), (steps[i + 1][0] - 1.15, 2.2))
    # Return arrow runs below the boxes: next week starts again at step ①.
    ax.annotate("", xy=(1.25, 1.55), xytext=(9.2, 1.55), arrowprops=dict(
        arrowstyle="-|>", color="#5B6573", lw=1.1, connectionstyle="arc3,rad=-0.12"))
    ax.text(5.2, 0.85, "进入下一周", ha="center", fontsize=8, color="#5B6573")
    ax.text(5.2, 0.3, "每周记录：准备达标率 · 匹配达成率 · 高风险人数 · 平均缺口 · 教师工时占用",
            ha="center", fontsize=9, color="#17324D")
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def fig_timeseries(path: Path, data: dict, metric: str, ylabel: str, scale: float) -> None:
    fig, ax = plt.subplots(figsize=(6.6, 3.4))
    for key, row in data["four_scenarios"].items():
        marker, line, color = STYLES[key]
        weeks = [w["week"] for w in row["weekly"]]
        values = [w[metric] * scale for w in row["weekly"]]
        ax.plot(weeks, values, linestyle=line, marker=marker, markersize=3.5, color=color,
                linewidth=1.4, label=LABELS[key])
    ax.set_xlabel("仿真周")
    ax.set_ylabel(ylabel)
    ax.grid(alpha=0.3)
    ax.legend(fontsize=8, frameon=False)
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def fig_endpoint(path: Path, data: dict) -> None:
    rows = list(data["four_scenarios"].values())
    names = [r["label"] for r in rows]
    fig, axes = plt.subplots(1, 2, figsize=(7.6, 3.2))
    for ax, key, title in ((axes[0], "matched", "终点匹配达成率 / %"),
                           (axes[1], "high_risk", "终点高风险人数 / 人")):
        means = [r[key]["mean"] for r in rows]
        errs = [[r[key]["mean"] - r[key]["ci_low"] for r in rows],
                [r[key]["ci_high"] - r[key]["mean"] for r in rows]]
        ax.bar(names, means, yerr=errs, capsize=4, color=["#9E9E9E", "#6BAED6", "#FC9272", "#74C476"],
               edgecolor="#333333", hatch=None)
        for i, value in enumerate(means):
            ax.text(i, value + errs[1][i] + (0.6 if key == "matched" else 0.2), f"{value:.1f}",
                    ha="center", fontsize=8)
        ax.set_title(title, fontsize=10)
        ax.tick_params(axis="x", labelsize=8)
        ax.grid(axis="y", alpha=0.3)
        if key == "matched":
            ax.set_ylim(40, max(means) + 6)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def fig_sensitivity(path: Path, data: dict) -> None:
    sweeps = [("supply_ratio", "岗位供给比例"), ("counselor_count", "教师数量"),
              ("learning_rate", "学习率")]
    fig, axes = plt.subplots(1, 3, figsize=(8.4, 3.0))
    for ax, (key, label) in zip(axes, sweeps, strict=True):
        points = data["sweeps"][key]
        xs = [p["value"] for p in points]
        for field, name, (marker, line, color) in (
            ("baseline_matched", "基线", STYLES["baseline"]),
            ("precision_matched", "精准帮扶", STYLES["precision_coaching"]),
        ):
            means = [p[field]["mean"] for p in points]
            lows = [p[field]["ci_low"] for p in points]
            highs = [p[field]["ci_high"] for p in points]
            ax.plot(xs, means, marker=marker, linestyle=line, color=color, label=name, linewidth=1.3)
            ax.fill_between(xs, lows, highs, color=color, alpha=0.15)
        ax.set_xlabel(label, fontsize=9)
        ax.set_ylabel("终点匹配达成率 / %", fontsize=8)
        ax.grid(alpha=0.3)
        ax.tick_params(labelsize=8)
    axes[0].legend(fontsize=8, frameon=False)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def build_figures(data: dict) -> dict[str, Path]:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    figs = {
        "architecture": FIG_DIR / "fig1-architecture.png",
        "weekly_flow": FIG_DIR / "fig2-weekly-flow.png",
        "matched": FIG_DIR / "fig3-matched-rate.png",
        "high_risk": FIG_DIR / "fig4-high-risk.png",
        "endpoint": FIG_DIR / "fig5-endpoint.png",
        "sensitivity": FIG_DIR / "fig6-sensitivity.png",
    }
    fig_architecture(figs["architecture"])
    fig_weekly_flow(figs["weekly_flow"])
    fig_timeseries(figs["matched"], data, "matched_rate", "岗位匹配达成率 / %", 100)
    fig_timeseries(figs["high_risk"], data, "high_risk_count", "高风险人数 / 人", 1)
    fig_endpoint(figs["endpoint"], data)
    fig_sensitivity(figs["sensitivity"], data)
    return figs

# ---------------------------------------------------------------- docx helpers
BODY_PT = 12  # 小四
TABLE_PT = 10.5  # 五号
H1_PT = 14  # 四号, matches the template's heading runs (sz=28)
H2_PT = 12


def _fonts(run, east="宋体", ascii_="Times New Roman", size=BODY_PT, bold=None):
    run.font.name = ascii_
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), ascii_)
    rfonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def _spacing(p, before=0, after=0, line=1.5):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def _new_par_before(anchor):
    p = OxmlElement("w:p")
    anchor._p.addprevious(p)
    from docx.text.paragraph import Paragraph

    return Paragraph(p, anchor._parent)


class Writer:
    """Insert generated blocks before an anchor paragraph inside the template body."""

    def __init__(self, doc, anchor, figs):
        self.doc = doc
        self.anchor = anchor
        self.figs = figs
        self.h1_proto = None  # template "引言" paragraph: Heading 1 + auto numbering
        self.h2_count = 0

    def par(self):
        return _new_par_before(self.anchor)

    def text(self, text, *, indent=True, align=None, size=BODY_PT, bold=None, east="宋体"):
        p = self.par()
        _spacing(p)
        if indent:
            p.paragraph_format.first_line_indent = Pt(size * 2)
        if align is not None:
            p.alignment = align
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        # Split on citation markers so they can be superscripted.
        for part in re.split(r"(\[\d+\](?:\[\d+\])*)", text):
            if not part:
                continue
            run = p.add_run(part)
            _fonts(run, east=east, size=size, bold=bold)
            if re.fullmatch(r"(\[\d+\])+", part):
                run.font.superscript = True
        return p

    def h1(self, text):
        new = copy.deepcopy(self.h1_proto)
        self.anchor._p.addprevious(new)
        from docx.text.paragraph import Paragraph

        p = Paragraph(new, self.anchor._parent)
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
        p.runs[0].text = text
        self.h2_count = 0
        self.h1_index = getattr(self, "h1_index", 0) + 1
        return p

    def h2(self, text):
        self.h2_count += 1
        p = self.par()
        p.style = self.doc.styles["Heading 2"]
        _spacing(p, before=6, after=6, line=1.5)
        run = p.add_run(f"{self.h1_index}.{self.h2_count}  {text}")
        _fonts(run, east="黑体", ascii_="Times New Roman", size=H2_PT, bold=False)
        return p

    def bullets(self, items):
        for item in items:
            p = self.text("● " + item, indent=False)
            p.paragraph_format.left_indent = Pt(BODY_PT * 2)
            p.paragraph_format.first_line_indent = Pt(-BODY_PT)

    def formula(self, text):
        p = self.text(text, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        for r in p.runs:
            r.font.italic = True

    def caption(self, text):
        p = self.text(text, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=TABLE_PT, east="黑体")
        _spacing(p, before=3, after=6, line=1.25)
        return p

    def figure(self, key, caption):
        p = self.par()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _spacing(p, before=6, line=1.0)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(self.figs[key]), width=Cm(15))
        self.caption(caption)

    def table(self, caption, header, rows):
        cap = self.caption(caption)
        cap.paragraph_format.keep_with_next = True
        tbl = self.doc.add_table(rows=len(rows) + 1, cols=len(header))
        tbl.style = self.doc.styles["Table Grid"]
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        self.anchor._p.addprevious(tbl._tbl)
        size = 9 if len(header) >= 6 else TABLE_PT
        # Column widths follow the longest cell text so short columns don't steal space.
        def width_of(text):
            return sum(2 if ord(ch) > 127 else 1 for ch in str(text))
        longest = [max(width_of(r[c]) for r in [header, *rows]) for c in range(len(header))]
        weights = [min(max(v, 9), 40) for v in longest]
        usable = 15.0
        widths = [Cm(usable * w / sum(weights)) for w in weights]
        for row in tbl.rows:
            for ci, cell in enumerate(row.cells):
                cell.width = widths[ci]
        for ci, value in enumerate(header):
            self._cell(tbl.rows[0].cells[ci], value, bold=True, shade=True, size=size)
        for ri, row in enumerate(rows, start=1):
            for ci, value in enumerate(row):
                left = len(header) <= 4 and ci > 0 and width_of(value) > 16
                self._cell(tbl.rows[ri].cells[ci], value, size=size, left=left)
        # Repeat header row across pages.
        trpr = tbl.rows[0]._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trpr.append(el)
        spacer = self.par()
        _spacing(spacer, line=1.0)

    @staticmethod
    def _cell(cell, value, bold=False, shade=False, size=TABLE_PT, left=False):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.CENTER
        _spacing(p, line=1.15)
        run = p.add_run(str(value))
        _fonts(run, size=size, bold=bold)
        if shade:
            tcpr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "D9E2F3")
            tcpr.append(shd)

def _fill_underlined(paragraph, value, width_spaces):
    """Put a value into the template's underlined blank, keeping the line length."""
    blank = paragraph.runs[-1]
    used = sum(2 if ord(ch) > 127 else 1 for ch in value)
    pad = max(2, width_spaces - used)
    blank.text = "  " + value + " " * (pad - 2) if value else blank.text


def fill_cover(doc):
    p = doc.paragraphs
    p[3].runs[0].text = "智能社会治理"  # "Xxxx" + "赛道"
    p[4].runs[0].text = "智能社会治理智能体"  # "Xxxx" + "方向作品设计方案"
    _fill_underlined(p[9], COVER["project"], 48)
    _fill_underlined(p[11], COVER["unit"], 38)
    # 申请人 / 联系电话 are personal details left for the team to fill in (p[10], p[12]).


def _remove(paragraph):
    paragraph._p.getparent().remove(paragraph._p)


def build_docx(data, figs, blocks):
    doc = Document(str(TEMPLATE))
    fill_cover(doc)
    paras = doc.paragraphs
    abstract_title, h1_proto = paras[18], paras[19]
    anchor = paras[27]  # page break before 参考文献
    ref_heading, ref_proto = paras[28], paras[29]
    trailing = paras[30:]

    w = Writer(doc, anchor, figs)
    w.h1_proto = copy.deepcopy(h1_proto._p)
    abstract_proto = copy.deepcopy(abstract_title._p)
    for para in paras[18:27]:
        _remove(para)

    from docx.text.paragraph import Paragraph

    for block in blocks:
        kind = block[0]
        if kind == "abstract_title":
            el = copy.deepcopy(abstract_proto)
            anchor._p.addprevious(el)
            title = Paragraph(el, anchor._parent)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title.paragraph_format.page_break_before = True
            for r in title.runs:
                r.font.bold = True
                _fonts(r, east="黑体", size=16, bold=True)
        elif kind == "keywords":
            pk = w.text(block[1], indent=False)
            pk.runs[0].font.bold = True
        elif kind == "h1":
            hp = w.h1(block[1])
            if block[1] == "引言":
                hp.paragraph_format.page_break_before = True
        elif kind == "h2":
            w.h2(block[1])
        elif kind == "p":
            w.text(block[1])
        elif kind == "bullets":
            w.bullets(block[1])
        elif kind == "formula":
            w.formula(block[1])
        elif kind == "figure":
            w.figure(block[1], block[2])
        elif kind == "table":
            w.table(block[1], block[2], block[3])
        elif kind == "refs":
            refs = block[1]
    # References go after the template's 参考文献 heading, using its paragraph format.
    after = ref_proto
    for i, ref in enumerate(refs, start=1):
        el = copy.deepcopy(ref_proto._p)
        after._p.addnext(el)
        rp = Paragraph(el, ref_proto._parent)
        rp.paragraph_format.left_indent = Pt(BODY_PT * 2)
        rp.paragraph_format.first_line_indent = Pt(-BODY_PT * 2)
        run = rp.add_run(f"[{i}] {ref}")
        _fonts(run, size=TABLE_PT)
        after = rp
    _remove(ref_proto)
    for para in trailing:  # trailing page break + empty paragraphs → blank last page
        _remove(para)
    assert ref_heading.text == "参考文献"
    doc.core_properties.title = COVER["project"] + " 作品设计方案"
    doc.save(str(OUTPUT))


def build_markdown(blocks):
    lines = [f"# {COVER['project']} 作品设计方案", "",
             f"> 中国仿真创新应用大赛 · {COVER['track']} · {COVER['direction']}",
             "> 本文件由 `tools/build_proposal_docx.py` 与 Word 版同步生成，数字来自 "
             "`simulation-sensitivity.json`。", ""]
    h1 = h2 = 0
    for block in blocks:
        kind = block[0]
        if kind == "abstract_title":
            lines += ["## 摘要", ""]
        elif kind == "keywords":
            lines += [f"**{block[1]}**", ""]
        elif kind == "h1":
            h1 += 1
            h2 = 0
            lines += [f"## {h1}. {block[1]}", ""]
        elif kind == "h2":
            h2 += 1
            lines += [f"### {h1}.{h2} {block[1]}", ""]
        elif kind == "p":
            lines += [block[1], ""]
        elif kind == "bullets":
            lines += [f"- {item}" for item in block[1]] + [""]
        elif kind == "formula":
            lines += [f"`{block[1]}`", ""]
        elif kind == "figure":
            lines += [f"![{block[2]}](figures/{FIG_FILES[block[1]]})", "", f"*{block[2]}*", ""]
        elif kind == "table":
            _, caption, header, rows = block
            lines += [f"**{caption}**", "", "| " + " | ".join(header) + " |",
                      "|" + "---|" * len(header)]
            lines += ["| " + " | ".join(str(c) for c in r) + " |" for r in rows] + [""]
        elif kind == "refs":
            lines += ["## 参考文献", ""] + [f"[{i}] {r}" for i, r in enumerate(block[1], 1)] + [""]
    MARKDOWN.write_text("\n".join(lines), encoding="utf-8")


FIG_FILES = {}


def build_registration_docx(texts: dict[str, str]) -> dict[str, int]:
    counts = {key: char_count(value) for key, value in texts.items()}
    for key, count in counts.items():
        if count > LIMITS[key]:
            raise SystemExit(f"报名表「{key}」{count} 字，超过 {LIMITS[key]} 字限制")
    doc = Document(str(REG_TEMPLATE))
    table = doc.tables[0]

    def put(row, col, value, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, size=9):
        p = table.rows[row].cells[col].paragraphs[0]
        p.alignment = align
        if indent:
            p.paragraph_format.first_line_indent = Pt(size * 2)
        _spacing(p)
        p.paragraph_format.line_spacing = Pt(size + 3)  # exact spacing keeps the form on one page
        run = p.add_run(value)
        _fonts(run, size=size)

    put(0, 3, COVER["project"], align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=12)
    put(1, 2, COVER["unit"], align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=12)
    put(13, 2, texts["intro"])
    put(14, 2, texts["innovation"])
    put(15, 2, texts["value"])
    # The template's empty first paragraph pushes the closing 说明 note onto page 2
    # once the three fields hold text; dropping it keeps the form on a single page.
    first = doc.paragraphs[0]
    if not first.text.strip():
        _remove(first)
    doc.save(str(REG_OUTPUT))

    names = {"intro": "作品内容简介", "innovation": "创新点", "value": "推广应用价值"}
    lines = ["# 报名表文案", "",
             "> 由 `tools/build_proposal_docx.py` 生成并已填入 `中国仿真创新应用大赛报名表-职途智航.docx`；"
             "数字来自 `simulation-sensitivity.json`，字数按去除空白后的全部字符统计。", ""]
    for key, label in names.items():
        lines += [f"## {label}（限 {LIMITS[key]} 字以内）", "", texts[key], ""]
    lines += ["## 字数统计", "", "| 栏目 | 限制 | 实际 |", "|:---|:---|:---|"]
    lines += [f"| {names[k]} | ≤{LIMITS[k]} | {counts[k]} |" for k in names] + [""]
    REG_MARKDOWN.write_text("\n".join(lines), encoding="utf-8")
    return counts


def main():
    data = json.loads(DATA.read_text(encoding="utf-8"))
    tests = ROOT / "backend" / "tests"
    data["test_count"] = sum(f.read_text(encoding="utf-8").count("\n    def test_")
                             for f in tests.glob("test_*.py"))
    figs = build_figures(data)
    FIG_FILES.update({k: v.name for k, v in figs.items()})
    blocks = build_blocks(data)
    build_docx(data, figs, blocks)
    build_markdown(blocks)
    counts = build_registration_docx(build_registration(data))
    print(f"wrote {OUTPUT}\nwrote {MARKDOWN}\nwrote {REG_OUTPUT}\nfigures in {FIG_DIR}")
    print("registration char counts:", counts)


if __name__ == "__main__":
    main()
